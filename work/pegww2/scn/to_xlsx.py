# -*- coding: utf-8 -*-
"""把番号展开相关 TSV 汇总成一册可查看的 Excel（多工作表 + 表头冻结 + 按处置配色 + 筛选）。"""
import io, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
RES = r"d:\DSH\军事游戏汉化\pegww2\research"
DER = os.path.join(RES, "_derived")          # 生成器产出（可复算，勿手改）
OUT = os.path.join(RES, "番号展开-对照表-20261020.xlsx")

FILL = {
    "译": "C6EFCE", "保留": "D9D9D9", "部分待判": "FFEB9C",
    "待判": "FFC7CE", "舰名待译": "BDD7EE", "已定词": "C6EFCE", "仍待判": "FFC7CE",
}
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)

def read_tsv(path):
    """返回 (表头, 数据行)；跳过 # 注释行。"""
    lines = []
    for l in io.open(path, encoding="utf-8"):
        s = l.rstrip("\r\n")
        if not s or s.startswith("#"): continue
        lines.append(s.split("\t"))
    return (lines[0], lines[1:]) if lines else ([], [])

def add_sheet(wb, title, header, rows, color_col=None, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(1, c); cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for r in rows:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", str(x)) for x in r])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(header)), max(1, ws.max_row))
    for c in range(1, len(header) + 1):
        w = widths[c-1] if widths else max(10, min(42, max([len(str(header[c-1]))] + [len(str(r[c-1])) for r in rows[:400]]) + 3))
        ws.column_dimensions[get_column_letter(c)].width = w
    if color_col:
        idx = color_col - 1
        for i, r in enumerate(rows, 2):
            f = FILL.get(str(r[idx]).strip())
            if f:
                ws.cell(i, color_col).fill = PatternFill("solid", fgColor=f)
    return ws

wb = Workbook(); wb.remove(wb.active)

tbl_h, tbl = read_tsv(os.path.join(DER, "_番号展开-对照表-20261020.tsv"))
sfx_h, sfx = read_tsv(os.path.join(DER, "_番号-字母后缀码-待判清单-20261020.tsv"))
opq_h, opq = read_tsv(os.path.join(DER, "_番号-有意保留代号清单-20261020.tsv"))
ship_h, ship = read_tsv(os.path.join(DER, "_番号-舰名碎片-并入清单-20261020.tsv"))
nat_h, nat = read_tsv(os.path.join(RES, "_场景国籍精确归属-20261020.tsv"))
tok_h, tok = read_tsv(os.path.join(DER, "_u-token-map-mined.tsv"))

from collections import Counter
dispo = Counter(r[tbl_h.index("处置")] for r in tbl)
nats = Counter(r[tbl_h.index("国别")] for r in tbl)
tot = len(tbl)
ov = [["番号展开对照表 · 总览（依据：.SCN 逐单位精确国籍）"] + [""] * 4,
      ["总计行数", tot, "", "", ""],
      ["", "", "", "", ""],
      ["处置", "行数", "占比", "", ""]]
for k, v in dispo.most_common():
    ov.append([k, v, "%.1f%%" % (100.0 * v / tot), "", ""])
ov += [["", "", "", "", ""], ["国别", "行数", "占比", "", ""]]
for k, v in nats.most_common():
    ov.append([k, v, "%.1f%%" % (100.0 * v / tot), "", ""])
ws = wb.create_sheet("总览")
for r in ov: ws.append(r)
ws.cell(1, 1).font = Font(bold=True, size=12, color="1F4E79")
ws.cell(4, 1).font = ws.cell(4, 2).font = Font(bold=True)
for rr in range(5, 5 + len(dispo)):
    f = FILL.get(ws.cell(rr, 1).value, None)
    if f: ws.cell(rr, 1).fill = PatternFill("solid", fgColor=f)
ws.cell(7 + len(dispo), 1).font = ws.cell(7 + len(dispo), 2).font = Font(bold=True)
for c, w in zip("ABCDE", (16, 12, 12, 20, 20)):
    ws.column_dimensions[c].width = w

add_sheet(wb, "番号对照表", tbl_h, tbl, color_col=tbl_h.index("处置") + 1,
          widths=[10, 8, 14, 10, 16, 30, 32, 24, 10, 8])
pend = [r for r in tbl if r[tbl_h.index("处置")] in ("待判", "部分待判")]
add_sheet(wb, "待判清单", tbl_h, pend, color_col=tbl_h.index("处置") + 1,
          widths=[10, 8, 14, 10, 16, 30, 32, 24, 10, 8])
add_sheet(wb, "字母后缀码", sfx_h, sfx, color_col=sfx_h.index("状态") + 1,
          widths=[8, 10, 8, 26, 12, 34, 20, 10])
add_sheet(wb, "保留代号", opq_h, opq, widths=[12, 8, 26, 12])
add_sheet(wb, "舰名碎片", ship_h, ship, widths=[10, 8, 18, 14, 12])
add_sheet(wb, "场景国籍归属", nat_h, nat, widths=[22, 8, 26, 26, 10, 8, 8])
add_sheet(wb, "挖-token映射", tok_h, tok, widths=[16, 22, 8, 8, 60])

# ── 2026-10-20 收口批次（未解 token 裁决 + 实落盘 + 冲突/未落登记） ──
for fn, title, widths, cc, ccv in (
        ("_番号收口-定词表-20261020.tsv", "收口-定词表", [16, 8, 26, 22, 44], 4, ("译", "保留")),
        ("_番号收口-落盘表-20261020.tsv", "收口-落盘表", [30, 30, 30], None, None),
        ("_番号收口-冲突表-20261020.tsv", "收口-冲突表", [30, 50], None, None),
        ("_番号收口-纯数字键未落-20261020.tsv", "收口-纯数字未落", [14, 20], None, None)):
    p = os.path.join(DER, fn)
    if not os.path.isfile(p):
        continue
    h, r = read_tsv(p)
    add_sheet(wb, title, h, r, widths=widths)
    if cc:
        ws = wb[title]
        for i, row in enumerate(r, 2):
            v = str(row[cc - 1]).strip()
            ws.cell(i, cc).fill = PatternFill(
                "solid", fgColor="C6EFCE" if v.startswith("译") else "D9D9D9")

wb.save(OUT)
print("已生成:", OUT)
print("工作表:", wb.sheetnames)
print("行数: 对照表%d / 待判%d / 后缀码%d / 保留%d / 舰名%d / 场景国籍%d / token映射%d"
      % (len(tbl), len(pend), len(sfx), len(opq), len(ship), len(nat), len(tok)))
